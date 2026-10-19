from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Plano_Preventivo_EEV1500.xlsx")
F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="D9E2F3")
INP = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
AMB = PatternFill("solid", fgColor="FFE699")
GRN = PatternFill("solid", fgColor="C6EFCE")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()


def title(ws, text, sub=None, width=10):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, size=14, bold=True, color="1F3A5F")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=F, size=9, italic=True, color="555555")


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR
        cell.alignment = CEN
        cell.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, r1, r2, c1, c2, align=WRAP):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.alignment = align
            if cell.font is None or cell.font.name != F:
                cell.font = Font(name=F, size=10)


def classe_fmt(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"A"'], fill=RED, font=Font(name=F, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"B"'], fill=AMB, font=Font(name=F, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"C"'], fill=GRN, font=Font(name=F, bold=True)))


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4


# ---------------------------------------------------------------- LEIA-ME
ws = wb.active
ws.title = "LEIA-ME"
title(ws, "Plano de Manutenção Preventiva — Empilhadeira Elétrica Vonder EEV 1500 (1,5 tf)",
      "Base: Manual de Instruções Vonder EEV 1500 (cód. 68.57.150.010). Itens marcados 'Boa prática' não constam no manual.")
rows = [
    ("COMO USAR", ""),
    ("1. Criticidade", "Aba 'Criticidade': cada componente recebe notas de 1 a 5 (Segurança, Qualidade, Disponibilidade, Probabilidade). O índice e a classe A/B/C são calculados por fórmula. Ajuste as notas à sua realidade (células amarelas) e o plano inteiro se reordena na prioridade."),
    ("2. Plano", "Aba 'Plano Preventivo': todas as atividades, com frequência, gatilho em horas, executor, critério de aceitação, EPI, ferramenta e referência no manual. Ordenado por frequência e, dentro dela, por criticidade."),
    ("3. Controle", "Aba 'Controle Horímetro': digite o horímetro atual e o horímetro/data da última execução de cada nível. A planilha mostra o que está VENCIDO, a PROGRAMAR ou OK (vale o que vencer primeiro: horas ou dias)."),
    ("4. Execução", "Imprima 'Checklist Diário' (operador, 1 folha por semana) e use 'Ordem de Serviço' para cada preventiva semanal em diante. Guarde os registros (NR-12 exige registro das manutenções)."),
    ("5. Treinamento", "Aba 'Treinamento': módulos, carga horária sugerida e matriz de habilitação por assistente. Ninguém executa atividade de um nível para o qual não está habilitado."),
    ("", ""),
    ("NÍVEIS DE EXECUTOR", ""),
    ("N1 — Operador", "Inspeção pré-uso e pós-uso, carga de bateria. Só verifica e reporta; não desmonta nada. Exige treinamento e cartão de identificação (NR-11, itens 11.1.5 e 11.1.6; cartão com validade de 1 ano)."),
    ("N2 — Assistente de manutenção", "Limpeza, aperto, lubrificação, nível de óleo e eletrólito, densidade, inspeções visuais e funcionais, carga equalizadora. Só após módulos M1 a M7 aprovados."),
    ("N3 — Técnico / Assistência Técnica Autorizada Vonder", "Ajuste de freio, contator, escovas, troca de óleo hidráulico e de engrenagem, válvula de segurança, controlador, rolamentos. Durante a garantia, OBRIGATORIAMENTE na Assistência Autorizada (ver abaixo)."),
    ("", ""),
    ("CLASSES DE CRITICIDADE", ""),
    ("A (vermelho)", "Falha pode ferir pessoas ou derrubar carga. Item NC = máquina BLOQUEADA até correção (etiqueta 'NÃO OPERE')."),
    ("B (amarelo)", "Falha para a máquina ou gera risco moderado. Corrigir em até 48 h; operar só se o líder liberar por escrito."),
    ("C (verde)", "Falha com baixo impacto. Corrigir na próxima preventiva programada."),
    ("", ""),
    ("GARANTIA (manual, p. 64)", "90 dias de garantia legal + 9 meses contratual (total 12 meses). Perde a garantia se o produto for aberto, alterado, ajustado ou consertado por pessoa não autorizada pela Vonder, ou por falta de manutenção. Locadoras: só 90 dias. Por isso, nos primeiros 12 meses, todas as tarefas N3 vão para a Assistência Autorizada; as tarefas N1/N2 (sem abrir componentes) ficam com a equipe interna e devem ser registradas (prova de manutenção). Contato Vonder: 0800 723 4762, opção 1; www.vonder.com.br."),
    ("", ""),
    ("PREMISSAS", ""),
    ("Regime", "Frequências do manual consideram 8 h/dia e 200 h/mês. Em 2 ou 3 turnos, siga o horímetro (aba Controle), não o calendário."),
    ("Bateria", "Plano considera a bateria indicada pelo manual: chumbo-ácido ventilada 24 V – 216 Ah (não acompanha a máquina). Se usar bateria selada ou de lítio, as tarefas de eletrólito/densidade/equalização não se aplicam: siga o manual da bateria."),
    ("Normas citadas", "NR-11 (transporte e movimentação de materiais), NR-12 (máquinas: manutenção por profissional capacitado, com bloqueio de energia e registro), NR-6 (EPI), NR-10 (somente para intervenção no lado de corrente alternada do carregador/tomada). ISO 3691 é citada pelo manual."),
    ("", ""),
    ("INCONSISTÊNCIAS DO MANUAL (confirmar com a Assistência Autorizada)", ""),
    ("Folga do freio", "Tabelas 3, 4 e 5 dizem 0,2–0,8 mm; a Fig. 5 mostra 0,20–0,25 mm; o texto do item 9.6 manda ajustar se passar de 0,5 mm e fala em '0,2 mm e 03 mm' (erro de digitação). Adotado no plano: faixa-alvo 0,2–0,3 mm após ajuste; acima de 0,5 mm = ajustar; acima de 0,8 mm = bloquear a máquina."),
    ("Óleo hidráulico", "Tabela 7 manda trocar a cada 1.200 h / 6 meses; Tabela 8 a cada 1.000 h. Adotado: 1.000 h ou 6 meses, o que vier primeiro."),
    ("Caixa de redução", "Tabela 7 fala em troca de ÓLEO de engrenagem (85W-90 GL-5, item 9.9); Tabela 8 fala em troca de GRAXA a cada 1.000 h. Confirmar qual lubrificante vai na sua unidade antes da primeira troca."),
    ("Motores", "A p. 3 descreve motor de tração AC (sem escovas), mas as Tabelas 5–8 mandam inspecionar escovas de carvão nos motores. Confirmar com a AT quais motores têm escovas (provavelmente o da bomba hidráulica)."),
    ("Regra 17 (p. 6)", "'Proibido operar com o garfo acima de 500 mm' — interpretado como: proibido DESLOCAR a máquina com o garfo elevado acima de 500 mm."),
    ("'Interruptor de polimento'", "Tabela 5, item 15: provável erro de tradução (interruptor de fim de curso/limite). Confirmar."),
    ("Substituto do óleo", "O manual aceita SAE 20W/20 no lugar do HLP68 'se for difícil comprar'. Não recomendado misturar óleos: use sempre HLP ISO VG 46 e só troque de tipo com o sistema drenado e com aval da AT."),
    ("Lacunas", "O manual não cita inspeção de correntes de elevação nem dos garfos (trincas/desgaste), essenciais num mastro de 5,6 m. Foram incluídas como 'Boa prática' (referência: ISO 5057, inspeção de garfos)."),
]
r = 4
for k, v in rows:
    ws.cell(row=r, column=1, value=k)
    ws.cell(row=r, column=2, value=v)
    if v == "" and k:
        ws.cell(row=r, column=1).font = Font(name=F, bold=True, color="FFFFFF")
        ws.cell(row=r, column=1).fill = HDR
        ws.cell(row=r, column=2).fill = HDR
    else:
        ws.cell(row=r, column=1).font = Font(name=F, bold=True, size=10)
        ws.cell(row=r, column=2).font = Font(name=F, size=10)
        ws.cell(row=r, column=1).alignment = WRAP
        ws.cell(row=r, column=2).alignment = WRAP
    r += 1
ws.cell(row=17, column=1).fill = RED
ws.cell(row=18, column=1).fill = AMB
ws.cell(row=19, column=1).fill = GRN
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 120
landscape(ws)

# ---------------------------------------------------------------- PARÂMETROS
wp = wb.create_sheet("Parâmetros")
title(wp, "Parâmetros técnicos e limites de aceitação", "Valores do manual Vonder EEV 1500. Valores calculados estão em fórmula.")
header(wp, 4, ["Grupo", "Parâmetro", "Valor", "Unidade", "Fonte / observação"], [16, 46, 14, 10, 70])
params = [
    ("Máquina", "Capacidade nominal (altura de elevação até 2.500 mm)", 1500, "kg", "Manual p. 3 e Tabela 2. Acima de 2.500 mm a capacidade cai: consultar Fig. 1 (diagrama elevação x carga)."),
    ("Máquina", "Centro de carga", 600, "mm", "Tabela 2"),
    ("Máquina", "Elevação máxima (h3)", 5600, "mm", "Tabela 2"),
    ("Máquina", "Altura de segurança do mastro (acima dela, velocidade reduz para ~3 km/h)", 1800, "mm", "Item 5.2.3"),
    ("Máquina", "Massa sem bateria", 1328, "kg", "Tabela 2 — usar no içamento"),
    ("Máquina", "Largura mínima de corredor (Ast)", 2450, "mm", "Tabela 2"),
    ("Máquina", "Raio de giro (Wa)", 1575, "mm", "Tabela 2"),
    ("Ambiente", "Temperatura ambiente permitida (mín.)", -25, "°C", "Manual p. 3"),
    ("Ambiente", "Temperatura ambiente permitida (máx.)", 40, "°C", "Manual p. 3; a 40 °C, umidade ≤ 50%"),
    ("Ambiente", "Altitude máxima", 1000, "m", "Manual p. 3"),
    ("Bateria", "Tensão nominal / capacidade", "24 V / 216 Ah", "", "Tabela 2 (chumbo-ácido, 12 células)"),
    ("Bateria", "Tensão mínima de operação (proibido operar abaixo)", 20.4, "V", "Regra 15 (p. 5) e item 6.3 (1,7 V/célula)"),
    ("Bateria", "Densidade do eletrólito com carga completa (25 °C)", 1.28, "g/ml", "± 0,005 — itens 6.2.6 e Tabela 4"),
    ("Bateria", "Densidade para interromper a descarga e recarregar", 1.17, "g/ml", "Item 6.3.1"),
    ("Bateria", "Nível do eletrólito acima da placa protetora (mín.)", 15, "mm", "Item 6.2.3 (faixa 15–25 mm)"),
    ("Bateria", "Nível do eletrólito acima da placa protetora (máx.)", 25, "mm", "Item 6.2.3"),
    ("Bateria", "Temperatura máxima do eletrólito durante a carga", 45, "°C", "Item 6.2.5 — acima disso, reduzir corrente 50% ou parar até 35 °C"),
    ("Bateria", "Tensão final de carga (fim do 2º estágio)", 31.2, "V", "Item 6.2.6 (2,6 V/célula)"),
    ("Bateria", "Duração aproximada da carga comum", 12, "h", "Item 6.3.2 — carregar 130–140% do descarregado"),
    ("Bateria", "Correção de densidade por temperatura", "D25 = Dt + 0,0007 × (t − 25)", "", "Item 6.2.4 — t em °C, Dt = densidade lida"),
    ("Carregador", "Saída do carregador automático", "24 V / 30 A", "", "Item 6.9"),
    ("Freio", "Folga do freio — faixa-alvo após ajuste (mín.)", 0.2, "mm", "Fig. 5 / item 9.6 (ver inconsistência no LEIA-ME)"),
    ("Freio", "Folga do freio — faixa-alvo após ajuste (máx.)", 0.3, "mm", "Adotado no plano"),
    ("Freio", "Folga do freio — acima deste valor, AJUSTAR", 0.5, "mm", "Item 9.6"),
    ("Freio", "Folga do freio — acima deste valor, BLOQUEAR a máquina", 0.8, "mm", "Tabelas 3, 4 e 5 (limite superior)"),
    ("Rodas", "Roda de direção/tração — diâmetro novo", 250, "mm", "Tabela 2 (Ø 250 x 70)"),
    ("Rodas", "Roda de direção/tração — diâmetro mínimo (−5%)", "=C30*0.95", "mm", "Tabela 3, item 8: substituir quando o diâmetro reduzir 5%"),
    ("Rodas", "Roda do garfo — diâmetro novo", 80, "mm", "Tabela 2 (Ø 80 x 84)"),
    ("Rodas", "Roda do garfo — diâmetro mínimo (−5%)", "=C32*0.95", "mm", "Tabela 3, item 8"),
    ("Hidráulico", "Volume de óleo hidráulico", 12, "L", "Item 9.3"),
    ("Hidráulico", "Pressão da válvula de segurança (1.500 kg)", 16, "MPa", "Item 12.1 — SOMENTE Assistência Autorizada ajusta"),
    ("Cabo", "Comprimento máximo do cabo da bateria externa (reboque)", 6, "m", "Item 4.1.2 — com bateria externa é proibido levantar carga"),
]
for i, p in enumerate(params, 5):
    for j, v in enumerate(p, 1):
        wp.cell(row=i, column=j, value=v)
body(wp, 5, 4 + len(params), 1, 5)
for i in range(5, 5 + len(params)):
    wp.cell(row=i, column=3).alignment = CEN
landscape(wp)

# ---------------------------------------------------------------- CRITICIDADE
wc = wb.create_sheet("Criticidade")
title(wc, "Matriz de criticidade",
      "IC = (2×S + Q + D) × P   |   Classe A: S = 5 ou IC ≥ 60   |   Classe B: S = 4 ou IC ≥ 30   |   Classe C: demais. Notas em amarelo são editáveis.")
wc["A3"] = "Escalas (1 a 5): S = gravidade para pessoas; Q = dano à carga/produto; D = impacto da parada; P = probabilidade de falha no seu uso."
wc["A3"].font = Font(name=F, size=9, italic=True)
header(wc, 5, ["ID", "Componente / sistema", "Modo de falha principal", "Consequência", "S", "Q", "D", "P", "IC", "Classe", "Estratégia de manutenção"],
       [6, 30, 40, 40, 5, 5, 5, 5, 7, 8, 44])
comps = [
    ("C01", "Freio eletromagnético", "Desgaste das placas de fricção / folga fora da faixa / bobina aberta", "Máquina não para; atropelamento, colisão, queda de carga", 5, 4, 4, 3, "Teste funcional diário; medição semanal da folga; limpeza trimestral; ajuste só por N3"),
    ("C02", "Botão de reversão de emergência (timão)", "Botão travado ou sem resposta", "Operador prensado entre máquina e parede/obstáculo", 5, 1, 2, 2, "Teste funcional diário (pré-uso)"),
    ("C03", "Parada de emergência / plugue / chave geral", "Não corta a energia", "Impossível parar a máquina em perda de controle", 5, 1, 3, 2, "Teste funcional diário; inspeção mensal do plugue"),
    ("C04", "Sistema hidráulico (bomba, cilindro, tubos, válvulas)", "Vazamento, queda de pressão, válvula de descida travada", "Garfo desce sozinho ou não desce; queda de carga; piso escorregadio", 5, 5, 4, 3, "Inspeção visual diária; nível diário; troca de óleo 1.000 h/6 meses; tubo de alta 2.000 h"),
    ("C05", "Mastro, roletes, correntes e proteção do mastro", "Rolete danificado, corrente alongada/sem lubrificação, mastro deformado", "Travamento ou queda da carga em altura (até 5,6 m)", 5, 5, 4, 2, "Visual diário; mensal; inspeção anual completa (boa prática)"),
    ("C06", "Garfos", "Trinca no calcanhar, deformação, desgaste", "Ruptura do garfo com carga", 5, 5, 3, 2, "Visual diário; inspeção anual detalhada (boa prática, ISO 5057)"),
    ("C07", "Contator principal e controle de elevação", "Contatos soldados ou carbonizados", "Perda de controle: garfo continua subindo após soltar o botão (item 5.3.1)", 5, 4, 5, 3, "Inspeção semanal/mensal; polimento trimestral; substituir por condição (N3)"),
    ("C08", "Bateria de tração 24 V – 216 Ah", "Eletrólito baixo, densidade baixa, vazamento de ácido, sulfatação", "Queimadura química, explosão por hidrogênio, parada", 4, 2, 5, 4, "Carga diária correta; eletrólito e densidade semanais; equalização a cada 2–3 meses"),
    ("C09", "Carregador e conexão de carga", "Conector danificado, aquecimento, curva de carga errada", "Incêndio, sobrecarga ou subcarga da bateria", 4, 1, 4, 3, "Visual diário; mensal"),
    ("C10", "Timão, direção e interruptor de trava (posições A/B/C)", "Trava não atua fora do ângulo de 45° ± 35°, folga, rolamento preso", "Máquina se desloca quando deveria estar travada; direção presa", 4, 2, 4, 3, "Teste diário; inspeção semanal e mensal"),
    ("C11", "Chicotes, fusíveis e terminais", "Terminal frouxo/oxidado, isolação danificada", "Curto-circuito, incêndio, parada", 4, 2, 5, 3, "Inspeção mensal; limpeza só com ar comprimido"),
    ("C12", "Fixadores estruturais (pinos e parafusos)", "Afrouxamento", "Soltura de roda ou componente; acidente", 4, 3, 3, 3, "Verificação diária visual; reaperto semanal"),
    ("C13", "Rodas e rolamentos (tração e garfo)", "Desgaste > 5% do diâmetro, rolamento danificado", "Instabilidade, tremor, dano ao piso e à carga", 3, 3, 4, 4, "Visual diário; medição mensal; substituição 1.200 h (Tabela 8)"),
    ("C14", "Motores (tração, bomba, direção)", "Escovas/rolamentos gastos, diodo/retificador com defeito", "Parada da máquina", 2, 2, 5, 3, "Inspeção mensal; trimestral/semestral por N3; rolamentos 1.000 h"),
    ("C15", "Caixa de redução", "Lubrificante degradado, ruído", "Parada da máquina", 2, 1, 5, 2, "Troca de lubrificante 1.000–1.200 h (confirmar tipo com AT)"),
    ("C16", "Controlador, painel e horímetro", "Código de falha, perda de parâmetros", "Parada da máquina", 3, 2, 5, 2, "Teste diário do painel; registrar códigos; diagnóstico por N3"),
    ("C17", "Buzina", "Não soa", "Pedestres não percebem a máquina", 3, 1, 1, 2, "Teste diário"),
    ("C18", "Sinalização, adesivos e placa de capacidade", "Ilegível ou ausente", "Sobrecarga e operação indevida", 4, 1, 1, 3, "Visual diário; repor imediatamente"),
]
for i, c in enumerate(comps, 6):
    for j, v in enumerate(c[:8], 1):
        wc.cell(row=i, column=j, value=v)
    wc.cell(row=i, column=9, value=f"=(2*E{i}+F{i}+G{i})*H{i}")
    wc.cell(row=i, column=10, value=f'=IF(OR(E{i}=5,I{i}>=60),"A",IF(OR(E{i}=4,I{i}>=30),"B","C"))')
    wc.cell(row=i, column=11, value=c[8])
last_c = 5 + len(comps)
body(wc, 6, last_c, 1, 11)
for i in range(6, last_c + 1):
    for col in (5, 6, 7, 8):
        wc.cell(row=i, column=col).fill = INP
        wc.cell(row=i, column=col).font = Font(name=F, size=10, color="0000FF")
    for col in (1, 5, 6, 7, 8, 9, 10):
        wc.cell(row=i, column=col).alignment = CEN
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", showErrorMessage=True, error="Use nota de 1 a 5")
wc.add_data_validation(dv)
dv.add(f"E6:H{last_c}")
classe_fmt(wc, f"J6:J{last_c}")
landscape(wc)


def classe_py(c):
    s, q, d, p = c[4:8]
    ic = (2 * s + q + d) * p
    return "A" if (s == 5 or ic >= 60) else ("B" if (s == 4 or ic >= 30) else "C")


CLS = {c[0]: classe_py(c) for c in comps}

# ---------------------------------------------------------------- PLANO
wpl = wb.create_sheet("Plano Preventivo")
title(wpl, "Plano preventivo — atividades, frequência, executor e critério de aceitação",
      "Classe puxada da aba Criticidade. Vale o que vencer primeiro: horas (horímetro) ou calendário. Ref. = item/tabela do manual.")
cols = ["Cód.", "Frequência", "Gatilho", "Comp.", "Classe", "Sistema", "Atividade", "Como fazer / critério de aceitação", "Executor", "EPI", "Ferramentas e materiais", "Ref. manual"]
header(wpl, 4, cols, [7, 13, 12, 7, 7, 16, 34, 62, 14, 22, 26, 14])

E_BAT = "Óculos de ampla visão ou protetor facial, luva nitrílica/PVC, avental de PVC, botina"
E_MEC = "Óculos, luva de vaqueta ou nitrílica, botina"
E_OP = "Botina, roupa justa (sem peças soltas)"
E_AR = "Óculos, protetor auricular, máscara PFF2, botina"

tasks = [
    # Diário pré-uso N1
    ("D01", "1-Diário (pré-uso)", "Cada turno", "C08", "Bateria", "Nível de carga no indicador", "Ligar chave geral, destravar trava elétrica e ler o medidor. Se a barra da extremidade esquerda estiver acesa: NÃO operar, carregar já. Nunca operar abaixo de 20,4 V.", "N1", E_OP, "—", "5.1.2 / Tab. 3-9"),
    ("D02", "1-Diário (pré-uso)", "Cada turno", "C08", "Bateria", "Cabo e conector da bateria", "Conector encaixado e firme; cabo sem corte, trinca ou marca de aquecimento. Qualquer dano = bloquear.", "N1", E_OP, "—", "Tab. 3-11"),
    ("D03", "1-Diário (pré-uso)", "Cada turno", "C08", "Bateria", "Vazamento de eletrólito (visual)", "Sem ácido derramado na caixa da bateria nem no piso. Não abrir tampas: só olhar. Vazamento = bloquear e chamar N2.", "N1", E_OP, "—", "Tab. 3-10 / 6.8"),
    ("D04", "1-Diário (pré-uso)", "Cada turno", "C04", "Hidráulico", "Vazamentos de óleo", "Olhar o piso sob a máquina, tubos, mangueiras e cilindro de elevação. Qualquer gotejamento = bloquear.", "N1", E_OP, "Lanterna", "5.1.1 / Tab. 3-4, 3-6"),
    ("D05", "1-Diário (pré-uso)", "Cada turno", "C04", "Hidráulico", "Nível do óleo hidráulico", "Com o garfo na posição mais baixa, nível dentro da marca do reservatório. Abaixo = chamar N2 (não completar).", "N1", E_OP, "Lanterna", "9.3 / Tab. 3-5"),
    ("D06", "1-Diário (pré-uso)", "Cada turno", "C01", "Freio", "Teste funcional do freio", "Timão nas posições A e C: máquina freada e travada. Em velocidade baixa, soltar o timão: deve parar em curta distância, sem arrastar. Falhou = BLOQUEAR.", "N1", E_OP, "—", "5.1.3 / Tab. 3-1, 3-2"),
    ("D07", "1-Diário (pré-uso)", "Cada turno", "C02", "Segurança", "Botão de reversão de emergência", "Em velocidade lenta, avançando em direção ao operador, pressionar o botão de reversão no topo do timão: máquina deve inverter e se afastar. Falhou = BLOQUEAR.", "N1", E_OP, "—", "5.1.3 / 5.3.3"),
    ("D08", "1-Diário (pré-uso)", "Cada turno", "C03", "Segurança", "Parada de emergência", "Acionar o botão/plugue de emergência: toda a energia deve ser cortada. Rearmar e confirmar religamento.", "N1", E_OP, "—", "4.1.1 / Fig. 4"),
    ("D09", "1-Diário (pré-uso)", "Cada turno", "C10", "Direção", "Timão e interruptor de trava", "Timão gira suave, volta sozinho para a vertical; só desloca na área B (45° ± 35°). Nas áreas A e C deve apenas subir/descer garfo. Ruído da trava presente ao passar de A para B.", "N1", E_OP, "—", "5.2.1 / Tab. 3-1, 3-3"),
    ("D10", "1-Diário (pré-uso)", "Cada turno", "C07", "Elevação", "Funções de elevação e descida", "Subir e descer o garfo sem carga até ~1 m: movimento suave, para ao soltar o botão, sem ruído anormal. Garfo que continua subindo = cortar energia (emergência) e BLOQUEAR.", "N1", E_OP, "—", "5.1.3 / 5.3.1 / Tab. 3-15"),
    ("D11", "1-Diário (pré-uso)", "Cada turno", "C05", "Mastro", "Mastro, correntes e proteção do mastro", "Proteção do mastro montada e íntegra; correntes sem elo travado ou trinca; sem objetos presos. Proteção ausente = proibido operar.", "N1", E_OP, "Lanterna", "3.8 / Boa prática"),
    ("D12", "1-Diário (pré-uso)", "Cada turno", "C06", "Garfos", "Garfos", "Sem trinca (principalmente na dobra/calcanhar), sem empeno, pontas alinhadas. Dúvida = bloquear e chamar N2.", "N1", E_OP, "Lanterna", "8.b / Boa prática"),
    ("D13", "1-Diário (pré-uso)", "Cada turno", "C12", "Estrutura", "Fixadores visíveis", "Nenhum pino ou parafuso faltando ou solto nas rodas e no chassi.", "N1", E_OP, "—", "Tab. 3-7"),
    ("D14", "1-Diário (pré-uso)", "Cada turno", "C13", "Rodas", "Rodas", "Sem pedaços faltando, sem fita/arame enrolado, sem achatamento visível. Máquina parada muito tempo pode ter roda marcada: some após rodar um pouco (item 4.1.6).", "N1", E_OP, "—", "Tab. 3-8"),
    ("D15", "1-Diário (pré-uso)", "Cada turno", "C16", "Painel", "Painel e chave", "Ligar a trava elétrica: painel acende e mostra estado normal, sem código de falha. Anotar qualquer código no checklist.", "N1", E_OP, "—", "Tab. 3-13, 3-14"),
    ("D16", "1-Diário (pré-uso)", "Cada turno", "C17", "Segurança", "Buzina", "Pressionar: deve soar alto e claro.", "N1", E_OP, "—", "Tab. 3-12"),
    ("D17", "1-Diário (pré-uso)", "Cada turno", "C18", "Segurança", "Adesivos e placa de capacidade", "Legíveis e no lugar (capacidade, advertências, pontos de içamento).", "N1", E_OP, "—", "8.a"),
    # Pós-uso
    ("P01", "2-Diário (pós-uso)", "Fim do turno", "C08", "Bateria", "Carga da bateria", "Estacionar em área ventilada, sem chamas, sem fumar. Conectar o carregador (24 V / 30 A). Não interromper a carga comum (~12 h). Não deixar a bateria descarregada parada. Nunca carregar ligando o plugue direto na rede CA.", "N1", E_BAT, "Carregador 24 V / 30 A", "6.1 / 6.3 / Regra 16"),
    ("P02", "2-Diário (pós-uso)", "Fim do turno", "C04", "Geral", "Estacionamento seguro", "Local plano e designado (nunca em rampa), garfo na posição mais baixa, timão na posição intermediária, chave desligada e retirada.", "N1", E_OP, "—", "3.1.1"),
    ("P03", "2-Diário (pós-uso)", "Fim do turno", "C11", "Geral", "Limpeza a seco", "Remover manchas e sujeira com pano seco. Parte elétrica só com ar comprimido. NUNCA usar água.", "N1", E_OP, "Pano, ar comprimido", "3.1.2 / 8"),
    ("P04", "2-Diário (pós-uso)", "Fim do turno", "C12", "Geral", "Registro de anomalias", "Anotar no checklist qualquer deformação, trinca, ruído, vazamento ou código de falha e avisar o líder.", "N1", "—", "Checklist diário", "8.b"),
    # Semanal 50 h N2
    ("S01", "3-Semanal", "50 h / 7 dias", "C01", "Freio", "Medição da folga do freio", "Máquina bloqueada (ver LOTO). Medir com calibrador de lâminas a folga entre placa do freio e aço magnético em 3 pontos. 0,2–0,5 mm = OK; >0,5 mm = solicitar ajuste N3; >0,8 mm = BLOQUEAR.", "N2 (N3 na garantia)", E_MEC, "Calibrador de lâminas 0,05–1,00 mm", "Tab. 4-3 / 9.6"),
    ("S02", "3-Semanal", "50 h / 7 dias", "C01", "Freio", "Limpeza da engrenagem de giro", "Remover sujeira e óleo acumulados na engrenagem de giro do freio/direção com pincel e pano.", "N2", E_MEC, "Pincel, pano, desengraxante não corrosivo", "Tab. 4-2"),
    ("S03", "3-Semanal", "50 h / 7 dias", "C10", "Direção", "Ruído da trava A/B", "Girar o timão alternando entre as áreas A e B: deve haver o ruído característico da trava em ambos os sentidos.", "N2", E_OP, "—", "Tab. 4-1"),
    ("S04", "3-Semanal", "50 h / 7 dias", "C08", "Bateria", "Nível do eletrólito", "Após a carga completa, abrir as tampas e conferir: líquido 15–25 mm acima da placa protetora. Se baixo, completar SÓ com água destilada, com funil plástico. Nunca completar com ácido.", "N2", E_BAT, "Funil plástico, água destilada, lanterna", "Tab. 4-4 / 6.2.3"),
    ("S05", "3-Semanal", "50 h / 7 dias", "C08", "Bateria", "Densidade do eletrólito", "Após a carga, medir em todas as 12 células com densímetro. Corrigir: D25 = Dt + 0,0007 × (t − 25). Esperado 1,280 ± 0,005. Célula com diferença grande = bateria 'atrasada': programar equalização.", "N2", E_BAT, "Densímetro de aspiração com termômetro", "Tab. 4-5 / 6.2.4"),
    ("S06", "3-Semanal", "50 h / 7 dias", "C08", "Bateria", "Limpeza da bateria", "Tampas fechadas. Limpar topo e terminais com pano umedecido em solução de bicarbonato de sódio e secar. Nenhum objeto metálico sobre a bateria.", "N2", E_BAT, "Pano, bicarbonato de sódio, escova não metálica", "Tab. 4-6 / 6.5.b"),
    ("S07", "3-Semanal", "50 h / 7 dias", "C07", "Elétrico", "Inspeção do contator", "Máquina bloqueada e capacitores descarregados. Inspecionar visualmente: sem marcas de arco, terminais firmes, acionamento sem travar.", "N2 (N3 na garantia)", E_MEC, "Lanterna, multímetro CAT III", "Tab. 4-7"),
    ("S08", "3-Semanal", "50 h / 7 dias", "C12", "Estrutura", "Reaperto de fixadores", "Conferir e reapertar parafusos e pinos acessíveis (rodas, timão, proteções). Sem torque especificado no manual: apertar firme, sem forçar.", "N2", E_MEC, "Jogo de chaves combinadas e soquetes", "9.4.b"),
    ("S09", "3-Semanal", "50 h / 7 dias", "C13", "Rodas", "Estado das rodas", "Inspecionar desgaste, trincas e corpos estranhos; limpar.", "N2", E_MEC, "Paquímetro", "9.4.c"),
    # Mensal 200 h
    ("M01", "4-Mensal", "200 h / 30 dias", "C01", "Freio", "Folga do freio e teste de frenagem", "Mesma medição de S01, registrar valor na OS para acompanhar a tendência de desgaste.", "N2 (N3 na garantia)", E_MEC, "Calibrador de lâminas", "Tab. 5-4"),
    ("M02", "4-Mensal", "200 h / 30 dias", "C07", "Elétrico", "Contator, chave, plugue, controlador, fim de curso", "Testar função de cada um. Contator: desempenho e contato. Plugue: sem dano. Registrar códigos do controlador.", "N2 (N3 na garantia)", E_MEC, "Multímetro CAT III", "Tab. 5-12 a 5-16"),
    ("M03", "4-Mensal", "200 h / 30 dias", "C04", "Hidráulico", "Tubos, cilindro e nível de óleo", "Sem vazamento em tubos e cilindro; nível correto com garfo abaixado. Completar só com HLP ISO VG 46 limpo.", "N2", E_MEC, "Óleo HLP 46, funil limpo", "Tab. 5-8 a 5-10"),
    ("M04", "4-Mensal", "200 h / 30 dias", "C05", "Mastro", "Mastro, roletes e correntes", "Garfo abaixado e calçado. Roletes girando livres, sem folga; correntes lubrificadas, sem elo duro, trinca ou alongamento visível; mastro sem deformação.", "N2", E_MEC, "Lanterna, lubrificante para correntes", "Boa prática"),
    ("M05", "4-Mensal", "200 h / 30 dias", "C10", "Direção", "Timão: estanqueidade, giro e trava", "Giro suave, sem folga; ruído da trava presente; quadro e prendedores sem trinca e apertados.", "N2", E_MEC, "Jogo de chaves", "Tab. 5-3, 5-5, 5-6"),
    ("M06", "4-Mensal", "200 h / 30 dias", "C13", "Rodas", "Haste de ligação e transportador de roda", "Sem trinca, flexão ou deformação; lubrificar pontos de graxa com graxa de lítio NLGI 3. Medir diâmetro das rodas (limites na aba Parâmetros).", "N2", E_MEC, "Pistola de graxa, graxa lítio NLGI 3, paquímetro", "Tab. 5-7"),
    ("M07", "4-Mensal", "200 h / 30 dias", "C11", "Elétrico", "Fusíveis, cablagem e terminais", "Bloqueado e capacitores descarregados. Fusíveis íntegros; cabos flexíveis, sem ressecamento ou corte; terminais firmes e sem oxidação.", "N2", E_MEC, "Multímetro CAT III, spray limpa-contatos", "Tab. 5-20, 5-21"),
    ("M08", "4-Mensal", "200 h / 30 dias", "C14", "Motores", "Escovas e diodos dos motores", "Verificar desgaste das escovas e estado dos diodos nos motores de tração, elevação e direção.", "N3", E_MEC, "Multímetro, ferramental da AT", "Tab. 5-17 a 5-19"),
    ("M09", "4-Mensal", "200 h / 30 dias", "C08", "Bateria", "Eletrólito completo", "Nível, densidade e limpeza conforme S04–S06; registrar densidade de cada célula.", "N2", E_BAT, "Densímetro, água destilada", "Tab. 5-11"),
    ("M10", "4-Mensal", "200 h / 30 dias", "C09", "Carregador", "Carregador e tomada", "Conector sem trinca nem aquecimento; cabo íntegro; área de carga ventilada, sinalizada, com extintor e sem material combustível.", "N2", E_MEC, "—", "2.1.4.c / 6.1"),
    # Bimestral/trimestral bateria
    ("B01", "5-Equalização", "A cada 2–3 meses", "C08", "Bateria", "Carga equalizadora", "Carregar a 4 A; ao atingir 31,2 V com bolhas, reduzir para 2 A. Com carga total: pausa 30 min, carregar 1 A por 1 h; repetir até borbulhar forte assim que o carregador liga. Área ventilada, sem chamas.", "N2", E_BAT, "Carregador ajustável, multímetro, densímetro", "6.3.3 / 6.4"),
    # Trimestral 600 h N3
    ("T01", "6-Trimestral", "600 h / 3 meses", "C07", "Elétrico", "Polimento dos contatos do contator", "Polir a superfície dos contatos com lixa fina; substituir se não funcionar bem.", "N3", E_MEC, "Lixa fina, multímetro", "Tab. 6"),
    ("T02", "6-Trimestral", "600 h / 3 meses", "C01", "Freio", "Limpeza e desgaste das placas de fricção", "Limpar poeira das placas de fricção; verificar desgaste; ajustar folga para 0,2–0,3 mm com os 3 parafusos uniformemente; testar com 24 V CC (deve emitir som claro).", "N3", E_MEC, "Calibrador de lâminas, ferramental AT", "Tab. 6 / 9.6"),
    ("T03", "6-Trimestral", "600 h / 3 meses", "C14", "Motores", "Escovas de carvão e retificador", "Verificar escovas e retificador de selênio; substituir se necessário.", "N3", E_MEC, "Ferramental AT", "Tab. 6"),
    # Consumíveis 1000 h
    ("K01", "7-Semestral / 1.000 h", "1.000 h ou 6 meses", "C04", "Hidráulico", "Troca do óleo hidráulico + limpeza do filtro", "Garfo totalmente abaixado (sistema despressurizado). Drenar, limpar filtro do reservatório e abastecer 12 L de HLP ISO VG 46 (DIN 51524-2). Óleo usado: descarte como resíduo perigoso.", "N3", E_MEC, "HLP ISO VG 46 (12 L), bandeja de coleta", "Tab. 7 / Tab. 8 / 9.9"),
    ("K02", "7-Semestral / 1.000 h", "1.000 h ou 6 meses", "C15", "Transmissão", "Lubrificante da caixa de redução", "Trocar o lubrificante da caixa de redução (óleo 85W-90 GL-5 ou graxa, conforme confirmação da AT).", "N3", E_MEC, "Óleo 85W-90 GL-5", "Tab. 7 / Tab. 8"),
    ("K03", "7-Semestral / 1.000 h", "1.000 h ou 6 meses", "C14", "Motores", "Rolamentos e escovas dos motores", "Motor de tração: rolamentos; motores de direção e bomba: escovas e rolamentos.", "N3", E_MEC, "Ferramental AT", "Tab. 8"),
    ("K04", "7-Semestral / 1.000 h", "1.200 h ou 6 meses", "C13", "Rodas", "Rodas e rolamentos do garfo", "Substituir conforme Tabela 8 (ou antes, se diâmetro < limite). Verificar desgaste da roda de tração.", "N3", E_MEC, "Rodas e rolamentos originais", "Tab. 7 / Tab. 8"),
    ("K05", "7-Semestral / 1.000 h", "1.200 h ou 6 meses", "C04", "Hidráulico", "Vedações do cilindro", "Verificar vazamento no cilindro de içamento; substituir vedações se houver dano.", "N3", E_MEC, "Kit de vedações original", "Tab. 7 / Tab. 8"),
    ("K06", "7-Semestral / 1.000 h", "1.200 h ou 6 meses", "C07", "Elétrico", "Contator, motor e freio (revisão semestral)", "Repetir T01, T02 e T03.", "N3", E_MEC, "Ferramental AT", "Tab. 7"),
    ("K07", "8-2.000 h", "2.000 h", "C04", "Hidráulico", "Tubo de óleo de alta pressão", "Substituir (ou antes, se houver dano).", "N3", E_MEC, "Tubo original", "Tab. 8"),
    # Anual boa prática
    ("A01", "9-Anual", "12 meses / 2.400 h", "C06", "Garfos", "Inspeção detalhada dos garfos", "Por pessoa qualificada: trincas (líquido penetrante no calcanhar), espessura, ângulo e alinhamento das pontas. Garfo com dano = substituir, nunca soldar.", "N3", E_MEC, "Paquímetro, esquadro, líquido penetrante", "Boa prática (ISO 5057)"),
    ("A02", "9-Anual", "12 meses / 2.400 h", "C05", "Mastro", "Correntes e mastro completo", "Medir alongamento das correntes, verificar pinos, ancoragens e roletes; lubrificar.", "N3", E_MEC, "Régua/gabarito de corrente", "Boa prática"),
    ("A03", "9-Anual", "12 meses / 2.400 h", "C04", "Hidráulico", "Pressão da válvula de segurança", "Somente AT: medir com manômetro ≥ 20 MPa. Esperado 16 MPa (1.500 kg). Usuário nunca ajusta.", "N3", E_MEC, "Manômetro ≥ 20 MPa", "12.1"),
    ("A04", "9-Anual", "12 meses / 2.400 h", "C08", "Bateria", "Avaliação da bateria", "Densidade de todas as células após carga completa e equalização; histórico de quedas de autonomia. Bateria tem vida útil limitada: planejar substituição.", "N2", E_BAT, "Densímetro, multímetro", "6.7 / Boa prática"),
]
order = {"A": 0, "B": 1, "C": 2}
tasks.sort(key=lambda t: (t[1], order[CLS[t[3]]], t[0]))
for i, t in enumerate(tasks, 5):
    vals = [t[0], t[1][2:], t[2], t[3], None, t[4], t[5], t[6], t[7], t[8], t[9], t[10]]
    for j, v in enumerate(vals, 1):
        wpl.cell(row=i, column=j, value=v)
    wpl.cell(row=i, column=5, value=f"=INDEX(Criticidade!$J$6:$J${last_c},MATCH(D{i},Criticidade!$A$6:$A${last_c},0))")
last_t = 4 + len(tasks)
body(wpl, 5, last_t, 1, 12)
for i in range(5, last_t + 1):
    for col in (1, 3, 4, 5, 9):
        wpl.cell(row=i, column=col).alignment = CEN
classe_fmt(wpl, f"E5:E{last_t}")
wpl.auto_filter.ref = f"A4:L{last_t}"
landscape(wpl)
wpl.print_title_rows = "4:4"

# ---------------------------------------------------------------- CONTROLE HORÍMETRO
wh = wb.create_sheet("Controle Horímetro")
title(wh, "Controle de vencimento das preventivas", "Edite só as células amarelas. Valores de exemplo em azul: substitua pelos reais. Vence pelo que ocorrer primeiro (horas ou dias).")
wh["A4"] = "Horímetro atual (h)"
wh["B4"] = 1180
wh["A5"] = "Data de hoje"
wh["B5"] = "=TODAY()"
wh["B5"].number_format = "DD/MM/YYYY"
for c in ("A4", "A5"):
    wh[c].font = Font(name=F, bold=True)
wh["B4"].fill = INP
wh["B4"].font = Font(name=F, color="0000FF", bold=True)
header(wh, 7, ["Nível", "Intervalo (h)", "Intervalo (dias)", "Último horímetro", "Última data", "Próximo horímetro", "Próxima data", "Faltam (h)", "Faltam (dias)", "Status", "Executor"],
       [26, 12, 12, 14, 13, 14, 13, 11, 11, 13, 20])
from datetime import date
lv = [
    ("Semanal (50 h)", 50, 7, 1150, date(2026, 10, 13), "N2"),
    ("Mensal (200 h)", 200, 30, 1000, date(2026, 9, 21), "N2 + N3"),
    ("Equalização da bateria", None, 75, None, date(2026, 8, 20), "N2"),
    ("Trimestral (600 h)", 600, 91, 600, date(2026, 7, 20), "N3"),
    ("Semestral / consumíveis (1.000 h)", 1000, 182, 0, date(2026, 5, 10), "N3"),
    ("Tubo de alta pressão (2.000 h)", 2000, None, 0, date(2026, 5, 10), "N3"),
    ("Anual (boa prática)", 2400, 365, 0, date(2026, 5, 10), "N3"),
]
for i, l in enumerate(lv, 8):
    wh.cell(row=i, column=1, value=l[0])
    wh.cell(row=i, column=2, value=l[1])
    wh.cell(row=i, column=3, value=l[2])
    wh.cell(row=i, column=4, value=l[3])
    wh.cell(row=i, column=5, value=l[4])
    wh.cell(row=i, column=6, value=f'=IF(OR(B{i}="",D{i}=""),"",D{i}+B{i})')
    wh.cell(row=i, column=7, value=f'=IF(OR(C{i}="",E{i}=""),"",E{i}+C{i})')
    wh.cell(row=i, column=8, value=f'=IF(F{i}="","",F{i}-$B$4)')
    wh.cell(row=i, column=9, value=f'=IF(G{i}="","",G{i}-$B$5)')
    wh.cell(row=i, column=10, value=(
        f'=IF(OR(AND(H{i}<>"",H{i}<=0),AND(I{i}<>"",I{i}<=0)),"VENCIDA",'
        f'IF(OR(AND(H{i}<>"",H{i}<=0.1*B{i}),AND(I{i}<>"",I{i}<=7)),"PROGRAMAR","OK"))'))
    wh.cell(row=i, column=11, value=l[5])
    for col in (4, 5):
        wh.cell(row=i, column=col).fill = INP
        wh.cell(row=i, column=col).font = Font(name=F, color="0000FF")
    wh.cell(row=i, column=5).number_format = "DD/MM/YYYY"
    wh.cell(row=i, column=7).number_format = "DD/MM/YYYY"
body(wh, 8, 7 + len(lv), 1, 11, CEN)
for i in range(8, 8 + len(lv)):
    for col in (4, 5):
        wh.cell(row=i, column=col).font = Font(name=F, color="0000FF")
rng = f"J8:J{7 + len(lv)}"
wh.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"VENCIDA"'], fill=RED, font=Font(name=F, bold=True)))
wh.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PROGRAMAR"'], fill=AMB, font=Font(name=F, bold=True)))
wh.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=GRN, font=Font(name=F, bold=True)))
wh["A16"] = "Regra: 'PROGRAMAR' aparece quando faltam ≤ 10% das horas ou ≤ 7 dias. Equalização: manual pede a cada 2–3 meses (75 dias adotado). Após executar, atualize 'Último horímetro' e 'Última data'."
wh["A16"].font = Font(name=F, size=9, italic=True)
landscape(wh)

# ---------------------------------------------------------------- CRONOGRAMA
wk = wb.create_sheet("Cronograma 52 sem")
title(wk, "Cronograma anual por calendário (referência para 1 turno de 8 h/dia)", "Informe a segunda-feira da semana 1 na célula amarela. Em mais de 1 turno, vale a aba Controle Horímetro.")
wk["A3"] = "Início (segunda-feira da semana 1):"
wk["A3"].font = Font(name=F, bold=True)
wk["C3"] = date(2026, 11, 2)
wk["C3"].fill = INP
wk["C3"].font = Font(name=F, color="0000FF", bold=True)
wk["C3"].number_format = "DD/MM/YYYY"
wk.column_dimensions["A"].width = 30
wk.column_dimensions["B"].width = 9
wk.cell(row=5, column=1, value="Nível")
wk.cell(row=5, column=2, value="A cada (sem.)")
wk.cell(row=6, column=1, value="Data início da semana")
for n in range(1, 53):
    col = 2 + n
    wk.cell(row=5, column=col, value=n)
    wk.cell(row=6, column=col, value=f"=$C$3+({get_column_letter(col)}5-1)*7")
    wk.cell(row=6, column=col).number_format = "DD/MM"
    wk.cell(row=6, column=col).alignment = Alignment(text_rotation=90, horizontal="center")
    wk.column_dimensions[get_column_letter(col)].width = 4
for col in range(1, 55):
    for rr in (5, 6):
        c = wk.cell(row=rr, column=col)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=9)
        c.fill = HDR
        c.border = BOX
        if rr == 5:
            c.alignment = CEN
wk.row_dimensions[6].height = 42
cr = [("Diário pré/pós-uso (N1)", 1), ("Semanal (N2)", 1), ("Mensal (N2/N3)", 4), ("Equalização bateria (N2)", 10),
      ("Trimestral (N3)", 13), ("Semestral / 1.000 h (N3)", 26), ("Anual – boa prática (N3)", 52)]
for i, (nm, every) in enumerate(cr, 7):
    wk.cell(row=i, column=1, value=nm)
    wk.cell(row=i, column=2, value=every)
    for n in range(1, 53):
        col = 2 + n
        wk.cell(row=i, column=col, value=f'=IF(MOD({get_column_letter(col)}$5,$B{i})=0,"●","")')
body(wk, 7, 6 + len(cr), 1, 54, CEN)
for i in range(7, 7 + len(cr)):
    wk.cell(row=i, column=1).alignment = Alignment(vertical="center")
    wk.cell(row=i, column=2).fill = INP
    wk.cell(row=i, column=2).font = Font(name=F, color="0000FF")
wk.conditional_formatting.add(f"C7:BB{6 + len(cr)}", FormulaRule(formula=['C7="●"'], fill=SUB))
wk.freeze_panes = "C7"
landscape(wk)

# ---------------------------------------------------------------- CHECKLIST DIÁRIO
wd = wb.create_sheet("Checklist Diário")
title(wd, "Checklist diário do operador — Empilhadeira EEV 1500", "Marque C (conforme), NC (não conforme) ou NA. Qualquer NC em item ★ = NÃO OPERE: coloque a etiqueta e avise o líder.")
for k, (lbl, col) in enumerate([("Semana de:", 1), ("Nº de série:", 4), ("Setor:", 7)]):
    wd.cell(row=3, column=col, value=lbl).font = Font(name=F, bold=True)
header(wd, 5, ["Item", "★", "Verificar", "Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"], [6, 4, 70, 7, 7, 7, 7, 7, 7, 7])
daily = [t for t in tasks if t[1].startswith("1-")]
for i, t in enumerate(daily, 6):
    wd.cell(row=i, column=1, value=t[0])
    wd.cell(row=i, column=2, value="★" if CLS[t[3]] == "A" else "")
    wd.cell(row=i, column=3, value=f"{t[5]}: {t[6].split('.')[0]}.")
end = 5 + len(daily)
extra = ["Horímetro no início", "Carga da bateria no início (%)", "Código de falha no painel (anotar)", "Operador (nome e assinatura)", "Líder: visto"]
for k, e in enumerate(extra, end + 1):
    wd.cell(row=k, column=3, value=e).font = Font(name=F, bold=True)
body(wd, 6, end + len(extra), 1, 10)
for i in range(6, end + len(extra) + 1):
    for col in (1, 2):
        wd.cell(row=i, column=col).alignment = CEN
    wd.row_dimensions[i].height = 30
dvc = DataValidation(type="list", formula1='"C,NC,NA"', allow_blank=True)
wd.add_data_validation(dvc)
dvc.add(f"D6:J{end}")
wd.conditional_formatting.add(f"D6:J{end}", CellIsRule(operator="equal", formula=['"NC"'], fill=RED, font=Font(name=F, bold=True)))
wd.cell(row=end + len(extra) + 2, column=1, value="Exemplo de preenchimento: D06 Seg = C; D10 Ter = NC → máquina bloqueada, etiqueta 'NÃO OPERE', OS aberta. Pós-uso: estacionar plano, garfo no chão, chave retirada, colocar para carregar em área ventilada.").font = Font(name=F, size=9, italic=True)
wd.page_setup.orientation = "portrait"
wd.page_setup.fitToWidth = 1
wd.page_setup.fitToHeight = 1
wd.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- ORDEM DE SERVIÇO
wo = wb.create_sheet("Ordem de Serviço")
title(wo, "Ordem de Serviço — Manutenção Preventiva", "Uma OS por intervenção. Arquivar por 5 anos (sugestão) como prova de manutenção para garantia e NR-12.")
for col, w in zip("ABCDEF", [24, 22, 22, 22, 22, 22]):
    wo.column_dimensions[col].width = w
fields = [("Nº OS:", "Data:", "Horímetro:"), ("Nível (Semanal/Mensal/...):", "Executor:", "Habilitação (N1/N2/N3):"),
          ("Hora início:", "Hora fim:", "Máquina bloqueada (LOTO) por:")]
r = 4
for f3 in fields:
    for k, f in enumerate(f3):
        wo.cell(row=r, column=1 + 2 * k, value=f).font = Font(name=F, bold=True, size=10)
        wo.cell(row=r, column=2 + 2 * k).border = BOX
    r += 1
r += 1
for k, h in enumerate(["Cód. tarefa", "Atividade", "Resultado (C/NC)", "Valor medido", "Ação tomada", "Visto"], 1):
    c = wo.cell(row=r, column=k, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = CEN
    c.border = BOX
ex = ("S01", "Medição da folga do freio", "C", "0,35 mm", "Nenhuma", "JS")
for k, v in enumerate(ex, 1):
    wo.cell(row=r + 1, column=k, value=v).font = Font(name=F, color="0000FF", italic=True)
body(wo, r + 1, r + 12, 1, 6)
for i in range(r + 1, r + 13):
    wo.row_dimensions[i].height = 22
r += 14
for k, h in enumerate(["Peça / consumível", "Código", "Quantidade", "Lote/NF", "", ""], 1):
    if h:
        c = wo.cell(row=r, column=k, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HDR
        c.border = BOX
body(wo, r + 1, r + 4, 1, 4)
r += 6
wo.cell(row=r, column=1, value="Anomalias encontradas / pendências:").font = Font(name=F, bold=True)
wo.merge_cells(start_row=r + 1, start_column=1, end_row=r + 3, end_column=6)
wo.cell(row=r + 1, column=1).border = BOX
r += 5
wo.cell(row=r, column=1, value="LIBERAÇÃO PARA USO").font = Font(name=F, bold=True, color="FFFFFF")
wo.cell(row=r, column=1).fill = HDR
lib = ["Ferramentas e materiais retirados da máquina", "Proteções e tampas recolocadas", "Teste funcional: freio, reversão de emergência, parada de emergência, elevação/descida, buzina",
       "Sem vazamentos após teste", "Bloqueio (cadeado/etiqueta) retirado pelo mesmo executor que bloqueou"]
for k, l in enumerate(lib, r + 1):
    wo.cell(row=k, column=1, value="☐ " + l).font = Font(name=F, size=10)
r += len(lib) + 2
wo.cell(row=r, column=1, value="Executor (assinatura):").font = Font(name=F, bold=True)
wo.cell(row=r, column=4, value="Líder / liberação (assinatura):").font = Font(name=F, bold=True)
wo.page_setup.orientation = "portrait"
wo.page_setup.fitToWidth = 1
wo.page_setup.fitToHeight = 1
wo.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- TREINAMENTO
wt = wb.create_sheet("Treinamento")
title(wt, "Programa de treinamento e matriz de habilitação", "Carga horária e nota mínima são sugestões (política interna). Reciclagem anual recomendada, alinhada à validade do cartão do operador (NR-11).")
header(wt, 4, ["Módulo", "Tema", "Público", "Carga (h)", "Teoria / prática", "Conteúdo-chave", "Avaliação para habilitar", "Base"],
       [8, 30, 18, 9, 14, 62, 36, 22])
mods = [
    ("M1", "Segurança e regras de operação", "Operadores e assistentes", 4, "Teoria", "Símbolos; regras do cap. 2 e 3; proibições (sobrecarga, passageiros, rampa, deslocar com garfo > 500 mm, piso irregular, ambiente inflamável); capacidade x altura (Fig. 1); centro de carga 600 mm.", "Prova escrita ≥ 80%", "Manual cap. 2–3; NR-11"),
    ("M2", "Inspeção pré-uso e operação", "Operadores e assistentes", 4, "Prática", "Checklist diário item a item com a máquina; teste de freio, reversão e parada de emergência; posições A/B/C do timão; empilhar/desempilhar (300 mm da prateleira, 100 mm acima, 200–300 mm do chão).", "3 inspeções completas supervisionadas sem erro", "Manual cap. 4, 5, 7"),
    ("M3", "Bateria: carga, eletrólito e emergência química", "Assistentes (operadores: só carga)", 3, "Teoria + prática", "Hidrogênio e ventilação; EPI de ácido; nível 15–25 mm; densímetro e correção por temperatura; carga comum e equalizadora; primeiros socorros (lavar com muita água; olhos: lavar e procurar médico); neutralização de derrame com cal ou bicarbonato.", "Medir 12 células e calcular D25 corretamente", "Manual cap. 6 e 11"),
    ("M4", "Bloqueio de energia e segurança na manutenção", "Assistentes", 2, "Teoria + prática", "Garfo no ponto mais baixo e calçado; chave e chave geral desligadas; desconectar bateria; cadeado + etiqueta; descarregar capacitores; sistema hidráulico despressurizado; trabalho em altura no mastro só com proteção; proibido fumar.", "Executar bloqueio completo sem falhas", "Manual 9.1; NR-12"),
    ("M5", "Manutenção N2", "Assistentes", 4, "Prática", "Tarefas semanais e mensais da aba Plano: limpeza a seco, reaperto, lubrificação (graxa lítio NLGI 3), nível de óleo HLP 46, medição de folga do freio, medição de rodas.", "Executar 1 preventiva semanal e 1 mensal supervisionadas", "Manual 9.4–9.6"),
    ("M6", "Diagnóstico básico e escalonamento", "Assistentes", 2, "Teoria", "Leitura de códigos e LEDs do controlador; tabela de falhas comuns (aba Diagnóstico); o que o assistente pode fazer e quando parar e chamar a AT.", "Estudo de 5 casos ≥ 80%", "Manual cap. 12–13"),
    ("M7", "Registro e documentação", "Assistentes e líderes", 1, "Prática", "Preencher checklist, OS e Controle Horímetro; etiqueta 'NÃO OPERE'; arquivo para garantia.", "Preencher 1 OS completa", "NR-12 (registro)"),
]
for i, m in enumerate(mods, 5):
    for j, v in enumerate(m, 1):
        wt.cell(row=i, column=j, value=v)
endm = 4 + len(mods)
body(wt, 5, endm, 1, 8)
wt.cell(row=endm + 1, column=3, value="Total (h)").font = Font(name=F, bold=True)
wt.cell(row=endm + 1, column=4, value=f"=SUM(D5:D{endm})").font = Font(name=F, bold=True)
r0 = endm + 4
wt.cell(row=r0 - 1, column=1, value="MATRIZ DE HABILITAÇÃO — preencha com a data de aprovação (dd/mm/aaaa). Vazio = não habilitado.").font = Font(name=F, bold=True)
hdrs = ["Nº", "Nome do assistente", "Função"] + [m[0] for m in mods] + ["Habilitado N2?", "Reciclagem em"]
for k, h in enumerate(hdrs, 1):
    c = wt.cell(row=r0, column=k, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = CEN
    c.border = BOX
people = [(1, "Exemplo: João Silva", "Assistente", [date(2026, 11, 3)] * 7)] + [(n, "", "", [None] * 7) for n in range(2, 9)]
first_mod_col = 4
last_mod_col = first_mod_col + len(mods) - 1
for k, (n, nm, fn, ds) in enumerate(people, r0 + 1):
    wt.cell(row=k, column=1, value=n)
    wt.cell(row=k, column=2, value=nm)
    wt.cell(row=k, column=3, value=fn)
    for j, d in enumerate(ds):
        c = wt.cell(row=k, column=first_mod_col + j, value=d)
        c.number_format = "DD/MM/YY"
        c.fill = INP
        if d:
            c.font = Font(name=F, color="0000FF")
    a, b = get_column_letter(first_mod_col), get_column_letter(last_mod_col)
    wt.cell(row=k, column=last_mod_col + 1, value=f'=IF(B{k}="","",IF(COUNT({a}{k}:{b}{k})={len(mods)},"SIM","NÃO"))')
    wt.cell(row=k, column=last_mod_col + 2, value=f'=IF(COUNT({a}{k}:{b}{k})=0,"",MAX({a}{k}:{b}{k})+365)')
    wt.cell(row=k, column=last_mod_col + 2).number_format = "DD/MM/YYYY"
body(wt, r0 + 1, r0 + len(people), 1, last_mod_col + 2, CEN)
for k in range(r0 + 1, r0 + len(people) + 1):
    for j in range(first_mod_col, last_mod_col + 1):
        wt.cell(row=k, column=j).fill = INP
wt.cell(row=r0 + 1, column=2).font = Font(name=F, color="0000FF", italic=True)
hc = get_column_letter(last_mod_col + 1)
wt.conditional_formatting.add(f"{hc}{r0+1}:{hc}{r0+len(people)}", CellIsRule(operator="equal", formula=['"SIM"'], fill=GRN))
wt.conditional_formatting.add(f"{hc}{r0+1}:{hc}{r0+len(people)}", CellIsRule(operator="equal", formula=['"NÃO"'], fill=RED))
landscape(wt)

# ---------------------------------------------------------------- DIAGNÓSTICO
wg = wb.create_sheet("Diagnóstico")
title(wg, "Falhas comuns: ação imediata e quem resolve", "Resumo das Tabelas 9 a 12 e do item 5.3 do manual. Na dúvida, bloquear e chamar a Assistência Autorizada.")
header(wg, 4, ["Sintoma", "Ação imediata (operador)", "Causas prováveis (manual)", "Quem resolve", "Ref."], [30, 42, 52, 18, 12])
diag = [
    ("Garfo continua subindo após soltar o botão", "Desligar a chave geral (cortar energia) imediatamente. Afastar pessoas. Não ficar sob o garfo. BLOQUEAR.", "Contator soldado / circuito de elevação", "N3", "5.3.1 / Tab. 12 cód. 3,8"),
    ("Freio não funciona em movimento", "Interromper a operação na hora; usar reversão/parada de emergência. BLOQUEAR.", "Parafusos do interruptor soltos; fio do freio solto; placas gastas", "N3", "5.3.2 / Tab. 9-4"),
    ("Máquina empurrando operador contra obstáculo", "Pressionar o botão de reversão de emergência no topo do timão.", "—", "N1 (ação) / N3 (se falhar)", "5.3.3"),
    ("Máquina não liga (contator não atua)", "Verificar carga da bateria, parada de emergência rearmada e conector. Persistindo, bloquear.", "Fusível queimado; chave de energia ou trava elétrica com defeito; conexão da bateria solta", "N2 (fusível, conexão) / N3", "Tab. 9-1"),
    ("Garfo não sobe", "Não insistir. Conferir carga x capacidade e bateria.", "Sobrecarga; óleo insuficiente; bateria fraca; bomba/motor; válvula de alívio", "N2 (nível, carga) / N3", "Tab. 9-7"),
    ("Garfo não desce após elevar", "Não ficar sob a carga. Isolar a área. BLOQUEAR.", "Mastro deformado; rolete danificado; orifício de retorno obstruído; válvula eletromagnética", "N3", "Tab. 9-8"),
    ("Só sobe/desce, não anda", "Conferir posição do timão (deve estar na área B).", "Timão fora da área B; contator; placa de circuito", "N1 / N3", "5.2.1 / Tab. 9-2"),
    ("Direção dura, ruído, motor esquentando", "Parar e reportar.", "Rolamento preso ou sem lubrificação; folga de engrenagem; rolamento da roda", "N3", "Tab. 9-5, 9-6"),
    ("Máquina treme em movimento", "Parar e reportar.", "Porcas da roda motriz soltas; rodas fora do mesmo plano", "N2 (reaperto) / N3", "Tab. 9-10"),
    ("Autonomia da bateria caiu", "Recarregar completamente; registrar.", "Célula danificada; eletrólito baixo ou contaminado", "N2 (eletrólito) / N3", "Tab. 9-9"),
    ("Derrame de ácido / contato com pele ou olhos", "Pele: lavar com muita água. Olhos: lavar com muita água e procurar médico. Roupa: retirar e lavar. Piso: neutralizar com cal ou bicarbonato e lavar.", "Tombamento ou dano da bateria", "Todos (1º socorro) / N2", "6.8"),
    ("LEDs do controlador: vermelho e amarelo piscando", "Desligar, anotar o código, não religar repetidamente.", "Controlador com defeito", "N3", "Tab. 11"),
    ("Código de falha no visor", "Anotar o código no checklist/OS e reportar.", "Ver Tabelas 10 e 12 do manual (códigos 1,2 a 9,4)", "N3", "Cap. 13"),
]
for i, d in enumerate(diag, 5):
    for j, v in enumerate(d, 1):
        wg.cell(row=i, column=j, value=v)
body(wg, 5, 4 + len(diag), 1, 5)
landscape(wg)

# ---------------------------------------------------------------- LUBRIFICANTES / FERRAMENTAS
wf = wb.create_sheet("Insumos e Ferramentas")
title(wf, "Lubrificantes, consumíveis, ferramentas e EPI recomendados", "Especificações do manual (item 9.9). Ferramentas especificadas por característica técnica, não por marca: compare orçamentos que atendam à especificação.")
header(wf, 4, ["Tipo", "Item", "Especificação mínima", "Uso / tarefa", "Fonte"], [16, 30, 52, 36, 20])
ins = [
    ("Lubrificante", "Óleo hidráulico (padrão)", "HLP ISO VG 46, DIN 51524-2 — 12 L por troca", "K01, M03", "Manual 9.9 / 9.3"),
    ("Lubrificante", "Óleo hidráulico (carga pesada, >60 °C)", "HLP ISO VG 68, DIN 51524-2", "Só se a AT recomendar", "Manual 9.9.b"),
    ("Lubrificante", "Óleo hidráulico (baixa temperatura, <60 °C)", "HLP ISO VG 32, DIN 51524-2", "Câmaras frias", "Manual 9.9.c"),
    ("Lubrificante", "Óleo de engrenagem", "SAE 85W-90, API GL-5 (hipoide)", "K02", "Manual 9.9"),
    ("Lubrificante", "Graxa", "Graxa de lítio NLGI 3", "M06, articulações", "Manual 9.9"),
    ("Bateria", "Água destilada / desmineralizada", "Pura, sem minerais; nunca água de torneira", "S04", "Manual 6.2.4"),
    ("Bateria", "Eletrólito (só carga inicial ou correção por profissional)", "Ácido sulfúrico 1,260 ± 0,005 (25 °C); correção com 1,40", "Carga inicial", "Manual 6.2.3 / 6.2.7"),
    ("Limpeza", "Spray limpa-contatos", "Para contatos elétricos, não condutivo", "M07, armazenamento", "Manual 10.1.h"),
    ("Limpeza", "Bicarbonato de sódio ou cal", "Neutralização de ácido", "S06, emergência", "Manual 6.8.d"),
    ("Ferramenta", "Calibrador de lâminas (cálibre de folga)", "Faixa 0,05–1,00 mm, lâminas de 0,05 mm", "S01, M01, T02", "Freio"),
    ("Ferramenta", "Densímetro de bateria", "De aspiração, escala 1,100–1,300 g/ml, com termômetro", "S05, M09, B01", "Manual 6.6.a"),
    ("Ferramenta", "Multímetro digital", "True RMS, CAT III 600 V (mín.), medição de tensão CC e continuidade", "S07, M02, M07, B01", "Boa prática"),
    ("Ferramenta", "Paquímetro", "0–300 mm, resolução 0,05 mm", "Rodas (limites 237,5 e 76 mm)", "Tab. 3-8"),
    ("Ferramenta", "Pistola de graxa", "Manual, com bico flexível", "M06", "—"),
    ("Ferramenta", "Funil e jarra plásticos", "Não metálicos, exclusivos para a bateria", "S04", "Manual 6.5.b"),
    ("Ferramenta", "Ar comprimido", "Bico de segurança com pressão reduzida; nunca contra o corpo", "P03, limpeza elétrica", "Manual 3.1.2"),
    ("Ferramenta", "Jogo de chaves combinadas e soquetes", "Métrico", "S08, M05", "—"),
    ("Segurança", "Kit de bloqueio (LOTO)", "Cadeado individual com chave única, etiqueta 'NÃO OPERE', garra para conector da bateria", "Toda manutenção N2/N3", "NR-12"),
    ("Segurança", "Calços / cavaletes de madeira", "Madeira quadrada resistente ao peso (1.328 kg + bateria)", "Garfo calçado, armazenamento", "Manual 10.1.d"),
    ("Segurança", "Extintor na área de carga", "Classe adequada a equipamentos elétricos (CO₂ ou pó ABC)", "Área de carga", "Manual 2.1.4.c"),
    ("Segurança", "Lava-olhos (de emergência ou frasco)", "Próximo à área de carga da bateria", "Emergência química", "Boa prática"),
    ("EPI", "Óculos de ampla visão / protetor facial", "Com CA válido", "Bateria, ar comprimido", "NR-6"),
    ("EPI", "Luvas resistentes a ácido", "Nitrílica ou PVC, com CA", "Bateria", "NR-6"),
    ("EPI", "Avental de PVC", "Com CA", "Bateria", "NR-6 / Manual cap. 11"),
    ("EPI", "Botina de segurança", "Com biqueira, com CA", "Todos", "Manual 2.1.2 / NR-6"),
]
for i, d in enumerate(ins, 5):
    for j, v in enumerate(d, 1):
        wf.cell(row=i, column=j, value=v)
body(wf, 5, 4 + len(ins), 1, 5)
landscape(wf)

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print("ok", len(tasks), "tarefas")
